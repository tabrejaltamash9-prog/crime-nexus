import asyncio
import logging
import json
import os
from pathlib import Path
from PIL import Image

# STUB mode flag to unblock hackathon development on machines without dependencies
USE_STUB = False

try:
    from pdf2image import convert_from_path
    import pdfplumber
except ImportError:
    logging.warning("PDF dependencies missing. Falling back to STUB mode.")
    USE_STUB = True

logger = logging.getLogger(__name__)

async def extract_text_from_document(file_path: Path, mime_type: str) -> str:
    """
    Extract text from PDF or images.
    Returns either a plain string (for simple extraction) or a JSON-serialized string 
    representing structured OCR blocks (when Surya OCR is used).
    """
    if USE_STUB:
        logger.info(f"STUB: Extracting text from {file_path.name}")
        await asyncio.sleep(2)  # Simulate processing time
        return f"--- Page 1 ---\n[STUB EXTRACTED TEXT for {file_path.name}]\nThis is simulated extracted content to unblock the NLP and RAG pipelines.\nPerson of interest: Rahul Kumar (Phone: 555-0199). Mention of 'R. Kumar' later."

    extracted_text = ""
    try:
        # Run actual extraction in a thread pool to avoid blocking async loop
        extracted_text = await asyncio.to_thread(_run_extraction, file_path, mime_type)
    except Exception as e:
        logger.error(f"OCR failed for {file_path.name}: {e}")
        extracted_text = ""
    
    return extracted_text

def _run_extraction(file_path: Path, mime_type: str) -> str:
    text = ""
    engine = os.environ.get("OCR_ENGINE", "surya").lower()
    
    if mime_type == "application/pdf":
        # 1. Try to extract selectable text
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text += f"--- Page {i+1} ---\n{page_text}\n"
        
        # 2. If no text was found, run OCR
        if not text.strip():
            logger.info(f"No selectable text found in PDF. Running OCR using {engine}...")
            images = convert_from_path(file_path)
            
            if engine == "surya":
                from services.ocr.surya_engine import extract_document_text
                from services.ocr.quality_gate import apply_quality_gate
                structured_pages = extract_document_text(images)
                quality_gated = apply_quality_gate(structured_pages)
                # Return the quality gated output as JSON string
                return json.dumps(quality_gated)
            else:
                logger.warning(f"Unsupported OCR engine: {engine}")
                return ""

    elif mime_type.startswith("image/"):
        logger.info(f"Running OCR on image using {engine}...")
        try:
            img = Image.open(file_path)
            if engine == "surya":
                from services.ocr.surya_engine import extract_document_text
                from services.ocr.quality_gate import apply_quality_gate
                structured_pages = extract_document_text([img])
                quality_gated = apply_quality_gate(structured_pages)
                return json.dumps(quality_gated)
            else:
                logger.warning(f"Unsupported OCR engine: {engine}")
                return ""
        except Exception as e:
            logger.error(f"Failed to process image {file_path}: {e}")
            return ""

    elif mime_type.startswith("text/") or file_path.suffix.lower() in ['.txt', '.md', '.json', '.csv', '.log', '.xml', '.py', '.yaml', '.yml']:
        logger.info("Extracting plain text...")
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
    elif "wordprocessingml.document" in mime_type:
        logger.info("Extracting word document text...")
        try:
            import docx
            doc = docx.Document(file_path)
            text = "\\n".join([para.text for para in doc.paragraphs])
        except ImportError:
            text = "[DOCX extraction requires python-docx package]"
    elif "presentationml.presentation" in mime_type:
        logger.info("Extracting powerpoint text...")
        try:
            import pptx
            ppt = pptx.Presentation(file_path)
            for slide in ppt.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\\n"
        except ImportError:
            text = "[PPTX extraction requires python-pptx package]"

    return text.strip()
